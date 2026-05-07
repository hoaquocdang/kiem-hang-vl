from __future__ import annotations

from datetime import datetime
from io import BytesIO

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import threading

from app import assets
from app.auth import (
    SESSION_COOKIE_NAME,
    SESSION_MAX_AGE_SECONDS,
    admin_setup_allowed,
    authenticate_user,
    cleanup_expired_sessions,
    count_other_active_admins,
    create_auth_session,
    create_user,
    delete_user,
    delete_auth_session,
    ensure_existing_admin_owner_machine,
    export_users_to_sheet,
    get_user_by_id,
    get_authenticated_user,
    get_session_token,
    list_users,
    reset_user_password,
    serialize_user,
    set_admin_owner_machine,
    set_user_active,
    sync_users_from_sheet,
    sync_users_from_sheet_if_needed,
    users_exist,
)
from app.config import ACCOUNT_SHEET_PATH, DATA_DIR, IS_CLOUD, IS_VERCEL, MUTATIONS_DISABLED, PUBLIC_DIR, STATIC_DIR
from app.sheets import enrich_item, prefetch_sheets
from app.db import (
    USERS_DATA_DIR,
    connection_scope,
    data_connection_scope,
    delete_user_database,
    init_db,
    migrate_legacy_stock_data,
    user_database_path,
)
from app.schemas import (
    CreateUserPayload,
    LoginPayload,
    ReasonPayload,
    ResetPasswordPayload,
    ScanPayload,
    SetupPayload,
    UserStatusPayload,
)
from app.services import parse_inventory_file, parse_scan_code, utc_now


app = FastAPI(title="Stock Audit App", version="0.3.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def add_private_network_access_header(request: Request, call_next):
    response = await call_next(request)
    if request.headers.get("origin"):
        response.headers["Access-Control-Allow-Private-Network"] = "true"
    return response


if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
else:
    @app.get("/static/styles.css")
    def fallback_styles() -> Response:
        return Response(assets.STYLES_CSS, media_type="text/css; charset=utf-8")

    @app.get("/static/app.js")
    def fallback_script() -> Response:
        return Response(assets.APP_JS, media_type="application/javascript; charset=utf-8")

    @app.get("/static/inventory-template.csv")
    def fallback_inventory_template() -> Response:
        headers = {"Content-Disposition": 'attachment; filename="inventory-template.csv"'}
        return Response(
            assets.INVENTORY_TEMPLATE_CSV,
            media_type="text/csv; charset=utf-8",
            headers=headers,
        )


VERCEL_DESKTOP_HTML = """
<!DOCTYPE html>
<html lang="vi">
  <head>
    <meta charset="UTF-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>Stock Audit App</title>
    <style>
      body { margin: 0; font-family: "Segoe UI", Arial, sans-serif; background: #f4f6f8; color: #17202a; }
      main { width: min(680px, calc(100% - 32px)); margin: 72px auto; padding: 28px; background: #fff; border: 1px solid #d8e0e8; border-radius: 8px; }
      h1 { margin-top: 0; font-size: 1.6rem; }
      p { line-height: 1.55; color: #617083; }
    </style>
  </head>
  <body>
    <main>
      <h1>Dùng Stock Audit App bản máy tính</h1>
      <p>Bản Vercel đã được chặn khởi tạo và ghi dữ liệu để tránh phát sinh lưu trữ trên server.</p>
      <p>Hãy mở file <strong>Stock Audit App.exe</strong> trên máy tính để đăng nhập, import tồn kho và scan hàng bằng dữ liệu local.</p>
    </main>
  </body>
</html>
"""


def _ensure_local_runtime() -> None:
    if not IS_VERCEL or IS_CLOUD:
        return
    raise HTTPException(
        status_code=503,
        detail=(
            "Bản Vercel đã bị khóa dữ liệu. Hãy dùng Stock Audit App bản máy tính "
            "để lưu dữ liệu local."
        ),
    )


def _ensure_mutations_allowed() -> None:
    if not MUTATIONS_DISABLED:
        return
    raise HTTPException(
        status_code=503,
        detail=(
            "Bản deploy trên Vercel không được ghi dữ liệu. "
            "Hãy dùng bản PC local để lưu phiên kiểm hàng trên máy tính."
        ),
    )


@app.on_event("startup")
def startup() -> None:
    if IS_VERCEL and not IS_CLOUD:
        return
    init_db()
    if not IS_CLOUD:
        with connection_scope() as connection:
            ensure_existing_admin_owner_machine(connection)
            sync_users_from_sheet(connection)
            admin = connection.execute(
                "SELECT id FROM app_users WHERE role = 'admin' ORDER BY id LIMIT 1"
            ).fetchone()
            if admin:
                migrate_legacy_stock_data(connection, admin["id"])
    threading.Thread(target=prefetch_sheets, daemon=True).start()


@app.get("/", response_model=None)
def index():
    index_path = PUBLIC_DIR / "index.html"
    if index_path.exists():
        return FileResponse(index_path)
    return HTMLResponse(assets.INDEX_HTML)


@app.post("/api/parse-inventory")
async def parse_inventory_stateless(file: UploadFile = File(...)) -> dict:
    """Stateless — parses inventory file and returns items as JSON. No auth, no storage."""
    items = await parse_inventory_file(file)
    return {"items": items, "count": len(items)}


@app.get("/health")
def health() -> dict:
    import os as _os
    db_url = _os.getenv("DATABASE_URL", "")
    db_pg_url = _os.getenv("DATABASE_POSTGRES_URL", "")
    db_url_unpool = _os.getenv("DATABASE_URL_UNPOOLED", "")
    return {
        "status": "ok",
        "timestamp": utc_now(),
        "local_bridge": not IS_VERCEL,
        "storage": {
            "data_dir": str(DATA_DIR),
            "users_data_dir": str(USERS_DATA_DIR),
            "is_vercel": IS_VERCEL,
            "is_cloud": IS_CLOUD,
            "writes_enabled": not MUTATIONS_DISABLED,
            "db_url_set": bool(db_url),
            "db_pg_url_set": bool(db_pg_url),
            "db_url_unpool_set": bool(db_url_unpool),
            "db_url_prefix": db_url[:20] if db_url else "",
        },
    }


def _session_response(payload: dict, token: str) -> JSONResponse:
    response = JSONResponse({**payload, "session_token": token})
    response.set_cookie(
        SESSION_COOKIE_NAME,
        token,
        max_age=SESSION_MAX_AGE_SECONDS,
        httponly=True,
        samesite="lax",
    )
    return response


def _clear_session_response(payload: dict) -> JSONResponse:
    response = JSONResponse(payload)
    response.delete_cookie(SESSION_COOKIE_NAME)
    return response


def _current_user(request: Request) -> dict:
    _ensure_local_runtime()
    with connection_scope() as connection:
        sync_users_from_sheet_if_needed(connection)
        cleanup_expired_sessions(connection)
        user = get_authenticated_user(connection, request)
    if not user:
        raise HTTPException(status_code=401, detail="Vui lòng đăng nhập để dùng app")
    return user


def _current_admin(user: dict = Depends(_current_user)) -> dict:
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Chỉ tài khoản admin được dùng chức năng này")
    return user


@app.get("/api/auth/status")
def auth_status(request: Request) -> dict:
    if IS_VERCEL and not IS_CLOUD:
        return {
            "authenticated": False,
            "setup_required": False,
            "desktop_required": True,
            "message": "Bản Vercel đã bị khóa dữ liệu. Hãy dùng Stock Audit App bản máy tính.",
        }

    with connection_scope() as connection:
        sync_users_from_sheet_if_needed(connection)
        cleanup_expired_sessions(connection)
        has_users = users_exist(connection)
        setup_required = not has_users and admin_setup_allowed()
        setup_locked = not has_users and not setup_required
        user = get_authenticated_user(connection, request)
    return {
        "authenticated": user is not None,
        "setup_required": setup_required,
        "setup_locked": setup_locked,
        "desktop_required": False,
        "message": (
            "Máy này chưa có danh sách tài khoản. Hãy dùng accounts.xlsx do admin cấp "
            "hoặc đăng nhập trên máy admin để tạo user."
            if setup_locked
            else ""
        ),
        "user": serialize_user(user) if user else None,
    }


@app.post("/api/auth/setup")
def setup_admin(payload: SetupPayload, request: Request) -> JSONResponse:
    _ensure_local_runtime()
    with connection_scope() as connection:
        sync_users_from_sheet(connection)
        if users_exist(connection):
            raise HTTPException(status_code=409, detail="Hệ thống đã có tài khoản admin")
        if not admin_setup_allowed():
            raise HTTPException(
                status_code=403,
                detail=(
                    "Máy này không được tạo admin mới. Hãy dùng accounts.xlsx do admin cấp "
                    "hoặc bật quyền setup trên máy admin."
                ),
            )
        user = create_user(
            connection,
            payload.username,
            payload.password,
            payload.display_name,
            role="admin",
            allow_admin_role=True,
        )
        set_admin_owner_machine(connection)
        export_users_to_sheet(connection)
        migrate_legacy_stock_data(connection, user["id"])
        token = create_auth_session(
            connection,
            user["id"],
            request.headers.get("user-agent", ""),
        )
    return _session_response(
        {
            "message": "Đã tạo tài khoản admin đầu tiên",
            "user": serialize_user(user),
        },
        token,
    )


@app.post("/api/auth/login")
def login(payload: LoginPayload, request: Request) -> JSONResponse:
    _ensure_local_runtime()
    with connection_scope() as connection:
        sync_users_from_sheet(connection)
        user = authenticate_user(connection, payload.username, payload.password)
        if not user:
            raise HTTPException(status_code=401, detail="Tên đăng nhập hoặc mật khẩu không đúng")
        token = create_auth_session(
            connection,
            user["id"],
            request.headers.get("user-agent", ""),
        )
    return _session_response(
        {
            "message": "Đăng nhập thành công",
            "user": serialize_user(user),
        },
        token,
    )


@app.post("/api/auth/logout")
def logout(request: Request) -> JSONResponse:
    _ensure_local_runtime()
    token = get_session_token(request)
    with connection_scope() as connection:
        delete_auth_session(connection, token)
    return _clear_session_response({"message": "Đã đăng xuất"})


def _serialize_admin_user(user: dict) -> dict:
    return serialize_user(user) | {
        "is_active": bool(user["is_active"]),
        "created_at": user.get("created_at") or "",
        "last_login_at": user.get("last_login_at") or "",
        "data_path": str(user_database_path(user["id"])),
    }


@app.get("/api/users")
def get_users(_: dict = Depends(_current_admin)) -> dict:
    with connection_scope() as connection:
        sync_users_from_sheet(connection)
        users = list_users(connection)
    return {
        "account_sheet_path": str(ACCOUNT_SHEET_PATH),
        "users": [_serialize_admin_user(user) for user in users],
    }


@app.post("/api/users")
def add_user(payload: CreateUserPayload, _: dict = Depends(_current_admin)) -> dict:
    with connection_scope() as connection:
        user = create_user(
            connection,
            payload.username,
            payload.password,
            payload.display_name,
            role=payload.role,
        )
        export_users_to_sheet(connection)
    with data_connection_scope(user["id"]):
        pass
    return {"message": "Đã tạo tài khoản", "user": _serialize_admin_user(user)}


@app.put("/api/users/{user_id}/password")
def admin_reset_password(
    user_id: int,
    payload: ResetPasswordPayload,
    request: Request,
    _: dict = Depends(_current_admin),
) -> dict:
    with connection_scope() as connection:
        user = reset_user_password(
            connection,
            user_id,
            payload.password,
            keep_session_token=get_session_token(request),
        )
        export_users_to_sheet(connection)
    return {"message": "Đã đặt lại mật khẩu", "user": _serialize_admin_user(user)}


@app.patch("/api/users/{user_id}/status")
def admin_update_user_status(
    user_id: int,
    payload: UserStatusPayload,
    admin: dict = Depends(_current_admin),
) -> dict:
    if user_id == admin["id"] and not payload.is_active:
        raise HTTPException(status_code=400, detail="Không thể khóa chính tài khoản đang đăng nhập")

    with connection_scope() as connection:
        target = get_user_by_id(connection, user_id)
        if not target:
            raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản")
        if target["role"] == "admin" and target["is_active"] and not payload.is_active:
            if count_other_active_admins(connection, user_id) <= 0:
                raise HTTPException(status_code=400, detail="Không thể khóa admin cuối cùng")
        user = set_user_active(connection, user_id, payload.is_active)
        export_users_to_sheet(connection)
    return {
        "message": "Đã mở khóa tài khoản" if payload.is_active else "Đã khóa tài khoản",
        "user": _serialize_admin_user(user),
    }


@app.delete("/api/users/{user_id}")
def admin_delete_user(user_id: int, admin: dict = Depends(_current_admin)) -> dict:
    if user_id == admin["id"]:
        raise HTTPException(status_code=400, detail="Không thể xóa chính tài khoản đang đăng nhập")

    with connection_scope() as connection:
        target = get_user_by_id(connection, user_id)
        if not target:
            raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản")
        if target["role"] == "admin" and target["is_active"]:
            if count_other_active_admins(connection, user_id) <= 0:
                raise HTTPException(status_code=400, detail="Không thể xóa admin cuối cùng")
        deleted = delete_user(connection, user_id)
        export_users_to_sheet(connection)
    delete_user_database(user_id)
    return {"message": f"Đã xóa tài khoản {deleted['username']}"}


def _get_active_session(connection) -> dict | None:
    user_id = getattr(connection, "current_user_id", None)
    if user_id is not None:
        return connection.execute(
            """
            SELECT id, name, source_filename, imported_at
            FROM stock_sessions
            WHERE is_active = 1 AND user_id = ?
            ORDER BY id DESC
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
    return connection.execute(
        """
        SELECT id, name, source_filename, imported_at
        FROM stock_sessions
        WHERE is_active = 1
        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()


_ITEM_SELECT_COLUMNS = """
    id,
    variant_key,
    COALESCE(product_name, '') AS product_name,
    COALESCE(operation_label, '') AS operation_label,
    COALESCE(operation_code, '') AS operation_code,
    COALESCE(operation_key, '') AS operation_key,
    COALESCE(price, '') AS price,
    COALESCE(color, '') AS color,
    COALESCE(size, '') AS size,
    COALESCE(source_note, '') AS source_note,
    COALESCE(color_name, '') AS color_name,
    COALESCE(material, '') AS material,
    COALESCE(style, '') AS style,
    COALESCE(form, '') AS form,
    COALESCE(attributes, '') AS attributes,
    stock_qty,
    scanned_qty,
    COALESCE(reason_code, '') AS reason_code,
    COALESCE(reason_note, '') AS reason_note,
    COALESCE(last_scanned_at, '') AS last_scanned_at
"""


def _build_item_snapshot(row: dict) -> dict:
    difference = row["scanned_qty"] - row["stock_qty"]
    remaining_qty = row["stock_qty"] - row["scanned_qty"]
    status = "matched"
    if difference > 0:
        status = "over"
    elif difference < 0:
        status = "short"
    return {**dict(row), "difference": difference, "remaining_qty": remaining_qty, "status": status}


def _build_summary(connection, session_id: int) -> dict:
    row = connection.execute(
        """
        SELECT
            COUNT(*) AS total_lines,
            COALESCE(SUM(stock_qty), 0) AS total_stock,
            COALESCE(SUM(scanned_qty), 0) AS total_scanned,
            COALESCE(SUM(CASE WHEN scanned_qty = stock_qty THEN 1 ELSE 0 END), 0) AS matched_lines,
            COALESCE(SUM(CASE WHEN scanned_qty > stock_qty THEN 1 ELSE 0 END), 0) AS over_lines,
            COALESCE(SUM(CASE WHEN scanned_qty < stock_qty THEN 1 ELSE 0 END), 0) AS short_lines
        FROM inventory_items
        WHERE session_id = ?
        """,
        (session_id,),
    ).fetchone() or {}
    total_stock = int(row.get("total_stock") or 0)
    total_scanned = int(row.get("total_scanned") or 0)
    return {
        "total_lines": int(row.get("total_lines") or 0),
        "total_stock": total_stock,
        "total_scanned": total_scanned,
        "matched_lines": int(row.get("matched_lines") or 0),
        "over_lines": int(row.get("over_lines") or 0),
        "short_lines": int(row.get("short_lines") or 0),
        "difference_total": total_scanned - total_stock,
    }


def _build_dashboard(connection, session_id: int) -> dict:
    rows = connection.execute(
        f"""
        SELECT {_ITEM_SELECT_COLUMNS}
        FROM inventory_items
        WHERE session_id = ?
            ORDER BY color, size, operation_code
        """,
        (session_id,),
    ).fetchall()
    items = [_build_item_snapshot(row) for row in rows]

    total_stock = sum(item["stock_qty"] for item in items)
    total_scanned = sum(item["scanned_qty"] for item in items)
    matched_lines = sum(1 for item in items if item["status"] == "matched")
    over_lines = sum(1 for item in items if item["status"] == "over")
    short_lines = sum(1 for item in items if item["status"] == "short")

    unmatched_events = connection.execute(
        """
        SELECT
            scan_code,
            COALESCE(color_code, '') AS color_code,
            COALESCE(size_code, '') AS size_code,
            COALESCE(operation_code, '') AS operation_code,
            COALESCE(operation_key, '') AS operation_key,
            quantity,
            note,
            created_at
        FROM scan_events
        WHERE session_id = ? AND status = 'unmatched'
        ORDER BY id DESC
        LIMIT 20
        """,
        (session_id,),
    ).fetchall()

    recent_events = connection.execute(
        """
        SELECT
            scan_code,
            COALESCE(color_code, '') AS color_code,
            COALESCE(size_code, '') AS size_code,
            COALESCE(operation_code, '') AS operation_code,
            COALESCE(operation_key, '') AS operation_key,
            quantity,
            status,
            COALESCE(note, '') AS note,
            created_at
        FROM scan_events
        WHERE session_id = ?
        ORDER BY id DESC
        LIMIT 20
        """,
        (session_id,),
    ).fetchall()

    return {
        "summary": {
            "total_lines": len(items),
            "total_stock": total_stock,
            "total_scanned": total_scanned,
            "matched_lines": matched_lines,
            "over_lines": over_lines,
            "short_lines": short_lines,
            "difference_total": total_scanned - total_stock,
        },
        "items": items,
        "discrepancies": [item for item in items if item["difference"] != 0 or item["reason_code"]],
        "unmatched_scans": unmatched_events,
        "recent_scans": recent_events,
    }


@app.get("/api/session")
@app.get("/session")
def get_current_session(user: dict = Depends(_current_user)) -> JSONResponse:
    with data_connection_scope(user["id"]) as connection:
        session = _get_active_session(connection)
        if not session:
            return JSONResponse(
                {
                    "session": None,
                    "summary": None,
                    "items": [],
                    "discrepancies": [],
                    "unmatched_scans": [],
                    "recent_scans": [],
                }
            )
        dashboard = _build_dashboard(connection, session["id"])
        return JSONResponse({"session": session, **dashboard})


@app.post("/api/session/import")
@app.post("/session/import")
async def import_inventory(
    file: UploadFile = File(...),
    session_name: str = Query(default="Phiên kiểm hàng mới"),
    user: dict = Depends(_current_user),
) -> dict:
    _ensure_mutations_allowed()
    items = await parse_inventory_file(file)
    for item in items:
        enriched = enrich_item(dict(item))
        item["color_name"] = enriched.get("color_name", "") or ""
        item["material"] = enriched.get("material", "") or ""
        item["style"] = enriched.get("style", "") or ""
        item["form"] = enriched.get("form", "") or ""
        item["attributes"] = enriched.get("attributes", "") or ""
    imported_at = utc_now()
    normalized_name = session_name.strip() or "Phiên kiểm hàng mới"

    with data_connection_scope(user["id"]) as connection:
        user_id = getattr(connection, "current_user_id", None)
        if user_id is not None:
            connection.execute(
                "UPDATE stock_sessions SET is_active = 0 WHERE is_active = 1 AND user_id = ?",
                (user_id,),
            )
            row = connection.execute(
                """
                INSERT INTO stock_sessions(name, source_filename, imported_at, is_active, user_id)
                VALUES (?, ?, ?, 1, ?)
                RETURNING id
                """,
                (normalized_name, file.filename or "inventory.xlsx", imported_at, user_id),
            ).fetchone()
        else:
            connection.execute(
                "UPDATE stock_sessions SET is_active = 0 WHERE is_active = 1"
            )
            row = connection.execute(
                """
                INSERT INTO stock_sessions(name, source_filename, imported_at, is_active)
                VALUES (?, ?, ?, 1)
                RETURNING id
                """,
                (normalized_name, file.filename or "inventory.xlsx", imported_at),
            ).fetchone()
        session_id = row["id"]
        connection.executemany(
            """
            INSERT INTO inventory_items(
                session_id,
                barcode,
                sku,
                variant_key,
                product_name,
                operation_label,
                operation_code,
                operation_key,
                price,
                color,
                size,
                source_note,
                color_name,
                material,
                style,
                form,
                attributes,
                stock_qty,
                scanned_qty
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
            """,
            [
                (
                    session_id,
                    item["variant_key"],
                    item["operation_label"],
                    item["variant_key"],
                    item["product_name"],
                    item["operation_label"],
                    item["operation_code"],
                    item["operation_key"],
                    item["price"],
                    item["color"],
                    item["size"],
                    item.get("source_note", ""),
                    item.get("color_name", ""),
                    item.get("material", ""),
                    item.get("style", ""),
                    item.get("form", ""),
                    item.get("attributes", ""),
                    item["stock_qty"],
                )
                for item in items
            ],
        )
        dashboard = _build_dashboard(connection, session_id)

    return {
        "message": f"Đã import {len(items)} biến thể tồn kho theo file chuẩn",
        "session": {
            "id": session_id,
            "name": normalized_name,
            "source_filename": file.filename or "inventory.xlsx",
            "imported_at": imported_at,
        },
        **dashboard,
    }


@app.post("/api/scan")
@app.post("/scan")
def scan_barcode(
    payload: ScanPayload,
    compact: bool = Query(default=False),
    user: dict = Depends(_current_user),
) -> dict:
    _ensure_mutations_allowed()
    now = utc_now()
    parsed = parse_scan_code(payload.barcode)

    with data_connection_scope(user["id"]) as connection:
        session = _get_active_session(connection)
        if not session:
            raise HTTPException(status_code=400, detail="Chưa có phiên kiểm hàng nào. Hãy import file tồn kho trước.")

        item = connection.execute(
            """
            SELECT id
            FROM inventory_items
            WHERE session_id = ?
              AND color = ?
              AND operation_key = ?
              AND (size = ? OR size = 'SALE')
            ORDER BY CASE WHEN size = ? THEN 0 WHEN size = 'SALE' THEN 1 ELSE 2 END,
                     id
            LIMIT 1
            """,
            (
                session["id"],
                parsed["color_code"],
                parsed["operation_key"],
                parsed["size_code"],
                parsed["size_code"],
            ),
        ).fetchone()

        if not item:
            note = (
                f"Không khớp file tồn theo màu {parsed['color_code']}, "
                f"size {parsed['size_code']}, tác nghiệp {parsed['operation_key']}"
            )
            connection.execute(
                """
                INSERT INTO scan_events(
                    session_id, item_id, barcode, scan_code, color_code, size_code, operation_code, operation_key,
                    quantity, status, note, created_at
                ) VALUES (?, NULL, ?, ?, ?, ?, ?, ?, ?, 'unmatched', ?, ?)
                """,
                (
                    session["id"],
                    parsed["raw_code"],
                    parsed["raw_code"],
                    parsed["color_code"],
                    parsed["size_code"],
                    parsed["operation_code"],
                    parsed["operation_key"],
                    payload.quantity,
                    note,
                    now,
                ),
            )
            scan_event = {
                "scan_code": parsed["raw_code"],
                "color_code": parsed["color_code"],
                "size_code": parsed["size_code"],
                "operation_code": parsed["operation_code"],
                "operation_key": parsed["operation_key"],
                "quantity": payload.quantity,
                "status": "unmatched",
                "note": note,
                "created_at": now,
            }
            if compact:
                summary = _build_summary(connection, session["id"])
                return {
                    "matched": False,
                    "message": note,
                    "scan_parts": parsed,
                    "session": session,
                    "summary": summary,
                    "scan_event": scan_event,
                    "compact": True,
                }
            dashboard = _build_dashboard(connection, session["id"])
            return {
                "matched": False,
                "message": note,
                "scan_parts": parsed,
                "session": session,
                **dashboard,
            }

        connection.execute(
            """
            UPDATE inventory_items
            SET scanned_qty = scanned_qty + ?, last_scanned_at = ?
            WHERE id = ?
            """,
            (payload.quantity, now, item["id"]),
        )
        connection.execute(
            """
            INSERT INTO scan_events(
                session_id, item_id, barcode, scan_code, color_code, size_code, operation_code, operation_key,
                quantity, status, note, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'matched', '', ?)
            """,
            (
                session["id"],
                item["id"],
                parsed["raw_code"],
                parsed["raw_code"],
                parsed["color_code"],
                parsed["size_code"],
                parsed["operation_code"],
                parsed["operation_key"],
                payload.quantity,
                now,
            ),
        )

        refreshed_item = connection.execute(
            f"SELECT {_ITEM_SELECT_COLUMNS} FROM inventory_items WHERE id = ?",
            (item["id"],),
        ).fetchone()
        snapshot = _build_item_snapshot(refreshed_item)

        if compact:
            summary = _build_summary(connection, session["id"])
            scan_event = {
                "scan_code": parsed["raw_code"],
                "color_code": parsed["color_code"],
                "size_code": parsed["size_code"],
                "operation_code": parsed["operation_code"],
                "operation_key": parsed["operation_key"],
                "quantity": payload.quantity,
                "status": "matched",
                "note": "",
                "created_at": now,
            }
            color_desc = snapshot.get("color_name", "") or snapshot["color"]
            material_desc = snapshot.get("material", "")
            extra_info = (
                f" ({color_desc}{' - ' + material_desc if material_desc else ''})"
                if snapshot.get("color_name") or material_desc
                else ""
            )
            return {
                "matched": True,
                "message": (
                    f"Đã ghi nhận {payload.quantity} cho màu {snapshot['color']}{extra_info} | "
                    f"size {snapshot['size']} | tác nghiệp {snapshot['operation_key'] or snapshot['operation_label']} | "
                    f"còn lại {snapshot['remaining_qty']}"
                ),
                "item": snapshot,
                "scan_parts": parsed,
                "session": session,
                "summary": summary,
                "scan_event": scan_event,
                "compact": True,
            }

        dashboard = _build_dashboard(connection, session["id"])

    color_desc = snapshot.get("color_name", "") or snapshot["color"]
    material_desc = snapshot.get("material", "")
    extra_info = f" ({color_desc}{' - ' + material_desc if material_desc else ''})" if snapshot.get("color_name") or material_desc else ""
    return {
        "matched": True,
        "message": (
            f"Đã ghi nhận {payload.quantity} cho màu {snapshot['color']}{extra_info} | "
            f"size {snapshot['size']} | tác nghiệp {snapshot['operation_key'] or snapshot['operation_label']} | "
            f"còn lại {snapshot['remaining_qty']}"
        ),
        "item": snapshot,
        "scan_parts": parsed,
        "session": session,
        **dashboard,
    }


@app.put("/api/items/{item_id}/reason")
@app.put("/items/{item_id}/reason")
def update_reason(
    item_id: int,
    payload: ReasonPayload,
    user: dict = Depends(_current_user),
) -> dict:
    _ensure_mutations_allowed()
    with data_connection_scope(user["id"]) as connection:
        session = _get_active_session(connection)
        if not session:
            raise HTTPException(status_code=400, detail="Chưa có phiên kiểm hàng nào")

        cursor = connection.execute(
            """
            UPDATE inventory_items
            SET reason_code = ?, reason_note = ?
            WHERE id = ? AND session_id = ?
            """,
            (payload.reason_code.strip(), payload.reason_note.strip(), item_id, session["id"]),
        )
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Không tìm thấy dòng hàng cần cập nhật")
        dashboard = _build_dashboard(connection, session["id"])
    return {"message": "Đã cập nhật lý do", "session": session, **dashboard}


@app.get("/api/export/discrepancies.xlsx")
@app.get("/export/discrepancies.xlsx")
def export_discrepancies(user: dict = Depends(_current_user)) -> StreamingResponse:
    with data_connection_scope(user["id"]) as connection:
        session = _get_active_session(connection)
        if not session:
            raise HTTPException(status_code=400, detail="Chưa có phiên kiểm hàng nào")
        dashboard = _build_dashboard(connection, session["id"])

    rows = [_report_row_from_item(item) for item in dashboard["discrepancies"]]
    workbook = _build_report_workbook(session, dashboard["summary"], rows)
    return _workbook_response(workbook)


@app.post("/api/export/discrepancies-stateless.xlsx")
async def export_discrepancies_stateless(request: Request) -> StreamingResponse:
    payload = await request.json()
    session = payload.get("session") if isinstance(payload.get("session"), dict) else {}
    summary = payload.get("summary") if isinstance(payload.get("summary"), dict) else {}
    rows = payload.get("rows") if isinstance(payload.get("rows"), list) else []
    workbook = _build_report_workbook(session, summary, rows)
    return _workbook_response(workbook)


def _report_row_from_item(item: dict) -> list:
    note = " | ".join(part for part in [item.get("source_note", ""), item["reason_note"]] if part)
    return [
        item["product_name"],
        item["operation_key"] or item["operation_label"],
        item["color"],
        item.get("color_name", ""),
        item.get("material", ""),
        item.get("style", ""),
        item["size"],
        item["stock_qty"],
        item["scanned_qty"],
        max(item["remaining_qty"], 0),
        max(item["difference"], 0),
        "Khớp" if item["status"] == "matched" else ("Dư" if item["status"] == "over" else "Thiếu"),
        item["reason_code"],
        note,
    ]


def _build_report_workbook(session: dict, summary: dict, rows: list[list]) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Báo cáo chênh lệch"
    worksheet.merge_cells("A1:N1")
    worksheet["A1"] = "BÁO CÁO CHÊNH LỆCH KIỂM HÀNG"
    worksheet["A1"].font = Font(bold=True, size=15, color="1D1D1F")
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet["A1"].fill = PatternFill("solid", fgColor="EAF3FF")
    worksheet["A2"] = "Phiên"
    worksheet["B2"] = session.get("name") or "-"
    worksheet["A3"] = "File tồn kho"
    worksheet["B3"] = session.get("source_filename") or "-"
    worksheet["A4"] = "Thời gian import"
    worksheet["B4"] = session.get("imported_at") or "-"
    worksheet["A6"] = "Tổng dòng"
    worksheet["B6"] = summary.get("total_lines", 0)
    worksheet["C6"] = "Tồn kho"
    worksheet["D6"] = summary.get("total_stock", 0)
    worksheet["E6"] = "Đã scan"
    worksheet["F6"] = summary.get("total_scanned", 0)
    worksheet["A7"] = "Khớp"
    worksheet["B7"] = summary.get("matched_lines", 0)
    worksheet["C7"] = "Dư"
    worksheet["D7"] = summary.get("over_lines", 0)
    worksheet["E7"] = "Thiếu"
    worksheet["F7"] = summary.get("short_lines", 0)
    headers = [
        "Tên hàng",
        "Tác nghiệp",
        "Mã màu",
        "Tên màu",
        "Chất liệu",
        "Kiểu dáng",
        "Size",
        "Tồn ban đầu",
        "Đã scan",
        "Còn lại",
        "Dư thừa",
        "Trạng thái",
        "Lý do",
        "Ghi chú",
    ]
    worksheet.append([])
    worksheet.append(headers)
    header_row = 9

    for row in rows:
        worksheet.append((row + [""] * len(headers))[: len(headers)])

    thin = Side(style="thin", color="D1D1D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1D1D1F")
    short_fill = PatternFill("solid", fgColor="FFF4E5")
    over_fill = PatternFill("solid", fgColor="FFECEC")
    matched_fill = PatternFill("solid", fgColor="EAF8EE")
    for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == header_row:
                cell.fill = header_fill
                cell.font = Font(bold=True, color="FFFFFF")
            elif cell.row > header_row:
                status = worksheet.cell(row=cell.row, column=12).value
                if status == "Dư":
                    cell.fill = over_fill
                elif status == "Thiếu":
                    cell.fill = short_fill
                elif status == "Khớp":
                    cell.fill = matched_fill

    for cell in ["A2", "A3", "A4", "A6", "C6", "E6", "A7", "C7", "E7"]:
        worksheet[cell].font = Font(bold=True)

    widths = [28, 12, 10, 14, 18, 14, 8, 12, 12, 12, 12, 12, 18, 34]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A10"
    worksheet.auto_filter.ref = f"A{header_row}:N{worksheet.max_row}"
    return workbook


def _workbook_response(workbook: Workbook) -> StreamingResponse:
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"bao-cao-chenh-lech-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.xlsx"
    download_headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=download_headers,
    )
