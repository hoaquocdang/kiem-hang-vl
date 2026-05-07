from __future__ import annotations

import hashlib
import hmac
import csv
import io
import platform
import re
import secrets
import time
import threading
from datetime import datetime, timedelta
from urllib.error import URLError
from urllib.request import Request as UrlRequest, urlopen

from fastapi import HTTPException, Request
from openpyxl import Workbook, load_workbook

from app.config import (
    ACCOUNT_SHEET_PATH,
    ACCOUNT_SHEET_URL,
    ACCOUNT_SHEET_URL_PATH,
    ADMIN_SETUP_ALLOWED,
    ADMIN_SETUP_UNLOCK_PATH,
    BUNDLED_ACCOUNT_SHEET_PATH,
    BUNDLED_ACCOUNT_SHEET_URL_PATH,
    IS_CLOUD,
)


SESSION_COOKIE_NAME = "stock_audit_session"
SESSION_MAX_AGE_SECONDS = 60 * 60 * 24 * 30
PASSWORD_HASH_ITERATIONS = 260_000
ACCOUNT_SHEET_NAME = "users"
ACCOUNT_META_SHEET_NAME = "_meta"
ACCOUNT_SHEET_HEADERS = [
    "username",
    "display_name",
    "role",
    "is_active",
    "password",
    "password_hash",
    "created_at",
    "last_login_at",
]
ADMIN_OWNER_MACHINE_KEY = "admin_owner_machine_id"
REMOTE_SHEET_SYNC_TIMEOUT_SECONDS = 6
SYNC_THROTTLE_SECONDS = 30
_sync_throttle_lock = threading.Lock()
_last_sync_at: float = 0.0


def auth_now() -> str:
    return datetime.utcnow().isoformat(timespec="seconds")


def normalize_username(username: str) -> str:
    return username.strip().lower()


def current_machine_id() -> str:
    system = platform.system().strip().lower() or "unknown"
    node = platform.node().strip().lower() or "unknown"
    return f"{system}:{node}"


def admin_setup_allowed() -> bool:
    if IS_CLOUD:
        return True  # Always allow setup; endpoint checks if users already exist
    return ADMIN_SETUP_ALLOWED or ADMIN_SETUP_UNLOCK_PATH.exists()


def get_auth_meta(connection, key: str) -> str | None:
    row = connection.execute("SELECT value FROM app_meta WHERE key = ?", (key,)).fetchone()
    return row["value"] if row else None


def set_auth_meta(connection, key: str, value: str) -> None:
    connection.execute(
        """
        INSERT INTO app_meta(key, value, updated_at)
        VALUES (?, ?, ?)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = excluded.updated_at
        """,
        (key, value, auth_now()),
    )


def get_admin_owner_machine_id(connection) -> str:
    return get_auth_meta(connection, ADMIN_OWNER_MACHINE_KEY) or ""


def set_admin_owner_machine(connection, machine_id: str | None = None) -> None:
    set_auth_meta(connection, ADMIN_OWNER_MACHINE_KEY, machine_id or current_machine_id())


def ensure_existing_admin_owner_machine(connection) -> None:
    if get_admin_owner_machine_id(connection):
        return
    if not admin_setup_allowed():
        return
    admin = connection.execute(
        "SELECT id FROM app_users WHERE role = 'admin' ORDER BY id LIMIT 1"
    ).fetchone()
    if admin:
        set_admin_owner_machine(connection)


def is_admin_owner_machine(connection) -> bool:
    if IS_CLOUD:
        return True  # No machine restriction in cloud mode
    owner_machine_id = get_admin_owner_machine_id(connection)
    return bool(owner_machine_id) and owner_machine_id == current_machine_id()


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        PASSWORD_HASH_ITERATIONS,
    )
    return f"pbkdf2_sha256${PASSWORD_HASH_ITERATIONS}${salt.hex()}${digest.hex()}"


def _looks_like_password_hash(value: str) -> bool:
    parts = value.split("$", 3)
    if len(parts) != 4:
        return False
    algorithm, iterations_text, salt_hex, digest_hex = parts
    if algorithm != "pbkdf2_sha256":
        return False
    try:
        int(iterations_text)
        bytes.fromhex(salt_hex)
        bytes.fromhex(digest_hex)
    except ValueError:
        return False
    return True


def verify_password(password: str, stored_hash: str) -> bool:
    try:
        algorithm, iterations_text, salt_hex, digest_hex = stored_hash.split("$", 3)
        iterations = int(iterations_text)
        if algorithm != "pbkdf2_sha256":
            return False
        expected = bytes.fromhex(digest_hex)
        actual = hashlib.pbkdf2_hmac(
            "sha256",
            password.encode("utf-8"),
            bytes.fromhex(salt_hex),
            iterations,
        )
    except (ValueError, TypeError):
        return False
    return hmac.compare_digest(actual, expected)


def serialize_user(user: dict) -> dict:
    return {
        "id": user["id"],
        "username": user["username"],
        "display_name": user["display_name"],
        "role": user["role"],
        "is_admin": user["role"] == "admin",
    }


def users_exist(connection) -> bool:
    row = connection.execute("SELECT 1 FROM app_users LIMIT 1").fetchone()
    return row is not None


def create_user(
    connection,
    username: str,
    password: str,
    display_name: str = "",
    role: str = "user",
    allow_admin_role: bool = False,
) -> dict:
    normalized_username = normalize_username(username)
    normalized_display_name = display_name.strip() or normalized_username
    if role not in {"admin", "user"}:
        raise HTTPException(status_code=400, detail="Vai trò tài khoản không hợp lệ")
    if role == "admin" and not allow_admin_role:
        raise HTTPException(status_code=403, detail="Không được tạo thêm tài khoản admin")

    try:
        row = connection.execute(
            """
            INSERT INTO app_users(username, display_name, password_hash, role, is_active, created_at)
            VALUES (?, ?, ?, ?, 1, ?)
            RETURNING id
            """,
            (
                normalized_username,
                normalized_display_name,
                hash_password(password),
                role,
                auth_now(),
            ),
        ).fetchone()
    except Exception as exc:
        message = str(exc).lower()
        if "unique" in message or "constraint" in message:
            raise HTTPException(status_code=409, detail="Tên đăng nhập đã tồn tại") from exc
        raise

    return connection.execute(
        """
        SELECT id, username, display_name, role, is_active, created_at, last_login_at
        FROM app_users
        WHERE id = ?
        """,
        (row["id"],),
    ).fetchone()


def list_users(connection) -> list[dict]:
    return connection.execute(
        """
        SELECT id, username, display_name, password_hash, role, is_active, created_at, last_login_at
        FROM app_users
        ORDER BY CASE WHEN role = 'admin' THEN 0 ELSE 1 END, username
        """
    ).fetchall()


def get_user_by_id(connection, user_id: int) -> dict | None:
    return connection.execute(
        """
        SELECT id, username, display_name, role, is_active, created_at, last_login_at
        FROM app_users
        WHERE id = ?
        """,
        (user_id,),
    ).fetchone()


def count_other_active_admins(connection, user_id: int) -> int:
    row = connection.execute(
        """
        SELECT COUNT(*) AS count
        FROM app_users
        WHERE role = 'admin'
          AND is_active = 1
          AND id != ?
        """,
        (user_id,),
    ).fetchone()
    return int(row["count"])


def reset_user_password(
    connection,
    user_id: int,
    new_password: str,
    keep_session_token: str = "",
) -> dict:
    user = get_user_by_id(connection, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản")
    connection.execute(
        "UPDATE app_users SET password_hash = ? WHERE id = ?",
        (hash_password(new_password), user_id),
    )
    if keep_session_token:
        connection.execute(
            "DELETE FROM auth_sessions WHERE user_id = ? AND token != ?",
            (user_id, keep_session_token),
        )
    else:
        delete_user_sessions(connection, user_id)
    return get_user_by_id(connection, user_id)


def set_user_active(connection, user_id: int, is_active: bool) -> dict:
    user = get_user_by_id(connection, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản")
    connection.execute(
        "UPDATE app_users SET is_active = ? WHERE id = ?",
        (1 if is_active else 0, user_id),
    )
    if not is_active:
        delete_user_sessions(connection, user_id)
    return get_user_by_id(connection, user_id)


def delete_user(connection, user_id: int) -> dict:
    user = get_user_by_id(connection, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản")
    delete_user_sessions(connection, user_id)
    connection.execute("DELETE FROM app_users WHERE id = ?", (user_id,))
    return user


def delete_user_sessions(connection, user_id: int) -> None:
    connection.execute("DELETE FROM auth_sessions WHERE user_id = ?", (user_id,))


def authenticate_user(connection, username: str, password: str) -> dict | None:
    user = connection.execute(
        """
        SELECT id, username, display_name, password_hash, role, is_active
        FROM app_users
        WHERE username = ?        """,
        (normalize_username(username),),
    ).fetchone()
    if not user or not user["is_active"]:
        return None
    if user["role"] == "admin" and not is_admin_owner_machine(connection):
        return None
    if not verify_password(password, user["password_hash"]):
        return None
    connection.execute(
        "UPDATE app_users SET last_login_at = ? WHERE id = ?",
        (auth_now(), user["id"]),
    )
    return user


def create_auth_session(connection, user_id: int, user_agent: str = "") -> str:
    token = secrets.token_urlsafe(32)
    expires_at = (datetime.utcnow() + timedelta(seconds=SESSION_MAX_AGE_SECONDS)).isoformat(
        timespec="seconds"
    )
    connection.execute(
        """
        INSERT INTO auth_sessions(token, user_id, created_at, expires_at, user_agent)
        VALUES (?, ?, ?, ?, ?)
        """,
        (token, user_id, auth_now(), expires_at, user_agent[:255]),
    )
    return token


def get_session_token(request: Request) -> str:
    authorization = request.headers.get("authorization", "")
    if authorization.lower().startswith("bearer "):
        return authorization.split(" ", 1)[1].strip()
    return request.cookies.get(SESSION_COOKIE_NAME, "")


def get_authenticated_user(connection, request: Request) -> dict | None:
    token = get_session_token(request)
    if not token:
        return None

    session = connection.execute(
        """
        SELECT
            u.id,
            u.username,
            u.display_name,
            u.role,
            u.is_active
        FROM auth_sessions s
        JOIN app_users u ON u.id = s.user_id
        WHERE s.token = ?
          AND s.expires_at > ?
        """,
        (token, auth_now()),
    ).fetchone()
    if not session or not session["is_active"]:
        return None
    if session["role"] == "admin" and not is_admin_owner_machine(connection):
        return None
    return session


def delete_auth_session(connection, token: str) -> None:
    if token:
        connection.execute("DELETE FROM auth_sessions WHERE token = ?", (token,))


def cleanup_expired_sessions(connection) -> None:
    connection.execute("DELETE FROM auth_sessions WHERE expires_at <= ?", (auth_now(),))


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_bool(value, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "active", "on", "đang dùng", "dang dung"}:
        return True
    if text in {"0", "false", "no", "n", "inactive", "off", "khóa", "khoa", "đã khóa", "da khoa"}:
        return False
    return default


def _normalize_sheet_role(value: str) -> str:
    role = value.strip().lower()
    return role if role in {"admin", "user"} else "user"


def _read_sheet_meta(workbook) -> dict[str, str]:
    if ACCOUNT_META_SHEET_NAME not in workbook.sheetnames:
        return {}
    sheet = workbook[ACCOUNT_META_SHEET_NAME]
    meta: dict[str, str] = {}
    for key, value in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
        key_text = _cell_text(key)
        if key_text:
            meta[key_text] = _cell_text(value)
    return meta


def _account_sheet_source_path():
    if not ACCOUNT_SHEET_PATH.exists():
        if BUNDLED_ACCOUNT_SHEET_PATH.exists():
            return BUNDLED_ACCOUNT_SHEET_PATH
        return ACCOUNT_SHEET_PATH
    if not BUNDLED_ACCOUNT_SHEET_PATH.exists():
        return ACCOUNT_SHEET_PATH
    try:
        if BUNDLED_ACCOUNT_SHEET_PATH.stat().st_mtime > ACCOUNT_SHEET_PATH.stat().st_mtime:
            return BUNDLED_ACCOUNT_SHEET_PATH
    except OSError:
        return ACCOUNT_SHEET_PATH
    return ACCOUNT_SHEET_PATH


def _text_file_value(path) -> str:
    try:
        if path.exists():
            return path.read_text(encoding="utf-8-sig").strip()
    except OSError:
        return ""
    return ""


def _account_sheet_url() -> str:
    if ACCOUNT_SHEET_URL:
        return ACCOUNT_SHEET_URL
    for path in [ACCOUNT_SHEET_URL_PATH, BUNDLED_ACCOUNT_SHEET_URL_PATH]:
        value = _text_file_value(path)
        if value:
            return value
    return ""


def _normalize_google_sheet_csv_url(url: str) -> str:
    spreadsheet_match = re.search(r"/spreadsheets/d/([^/?#]+)", url)
    if not spreadsheet_match:
        return url
    gid_match = re.search(r"[?&#]gid=(\d+)", url)
    gid = gid_match.group(1) if gid_match else "0"
    spreadsheet_id = spreadsheet_match.group(1)
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=csv&gid={gid}"


def _fetch_remote_account_rows(url: str) -> list[dict]:
    request = UrlRequest(
        _normalize_google_sheet_csv_url(url),
        headers={"User-Agent": "StockAuditApp/1.0"},
    )
    with urlopen(request, timeout=REMOTE_SHEET_SYNC_TIMEOUT_SECONDS) as response:
        content = response.read()
    text = content.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def _xlsx_account_rows_and_meta(path):
    workbook = load_workbook(path)
    sheet_meta = _read_sheet_meta(workbook)
    if ACCOUNT_SHEET_NAME not in workbook.sheetnames:
        return [], sheet_meta

    sheet = workbook[ACCOUNT_SHEET_NAME]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return [], sheet_meta
    headers = [_cell_text(value).lower() for value in header_row]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(
            {
                headers[index]: row[index] if index < len(row) else None
                for index in range(len(headers))
            }
        )
    return rows, sheet_meta


def _apply_account_rows(connection, rows: list[dict], allow_admin_roles: bool = True) -> tuple[bool, bool]:
    changed = False
    should_export = False

    for raw_values in rows:
        values = {
            _cell_text(key).lower(): value
            for key, value in raw_values.items()
            if _cell_text(key)
        }
        username = normalize_username(_cell_text(values.get("username")))
        if not username:
            continue

        display_name = _cell_text(values.get("display_name")) or username
        role = _normalize_sheet_role(_cell_text(values.get("role")) or "user")
        if role == "admin" and not (
            allow_admin_roles and (is_admin_owner_machine(connection) or admin_setup_allowed())
        ):
            continue

        is_active = _cell_bool(values.get("is_active"), True)
        plain_password = _cell_text(values.get("password"))
        sheet_password_hash = _cell_text(values.get("password_hash"))

        existing = connection.execute(
            """
            SELECT id, username, display_name, password_hash, role, is_active
            FROM app_users
            WHERE username = ?            """,
            (username,),
        ).fetchone()

        password_hash = ""
        if plain_password:
            should_export = True
            if not existing or not verify_password(plain_password, existing["password_hash"]):
                password_hash = hash_password(plain_password)
        elif _looks_like_password_hash(sheet_password_hash):
            password_hash = sheet_password_hash

        if existing:
            assignments = []
            params = []
            if existing["display_name"] != display_name:
                assignments.append("display_name = ?")
                params.append(display_name)
            if existing["role"] != role:
                assignments.append("role = ?")
                params.append(role)
            if bool(existing["is_active"]) != is_active:
                assignments.append("is_active = ?")
                params.append(1 if is_active else 0)
            if password_hash and existing["password_hash"] != password_hash:
                assignments.append("password_hash = ?")
                params.append(password_hash)
                delete_user_sessions(connection, existing["id"])
            if assignments:
                params.append(existing["id"])
                connection.execute(
                    f"UPDATE app_users SET {', '.join(assignments)} WHERE id = ?",
                    params,
                )
                changed = True
            continue

        if not password_hash:
            continue

        connection.execute(
            """
            INSERT INTO app_users(username, display_name, password_hash, role, is_active, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                username,
                display_name,
                password_hash,
                role,
                1 if is_active else 0,
                _cell_text(values.get("created_at")) or auth_now(),
            ),
        )
        changed = True

    return changed, should_export


def export_users_to_sheet(connection, raise_on_error: bool = True) -> None:
    if IS_CLOUD:
        return  # No persistent local filesystem in cloud mode
    ACCOUNT_SHEET_PATH.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    users_sheet = workbook.active
    users_sheet.title = ACCOUNT_SHEET_NAME
    users_sheet.append(ACCOUNT_SHEET_HEADERS)
    users_sheet.freeze_panes = "A2"

    for user in list_users(connection):
        users_sheet.append(
            [
                user["username"],
                user["display_name"],
                user["role"],
                "TRUE" if user["is_active"] else "FALSE",
                "",
                user.get("password_hash") or "",
                user.get("created_at") or "",
                user.get("last_login_at") or "",
            ]
        )

    for column_cells in users_sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        users_sheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 12),
            48,
        )

    meta_sheet = workbook.create_sheet(ACCOUNT_META_SHEET_NAME)
    meta_sheet.sheet_state = "hidden"
    meta_sheet.append([ADMIN_OWNER_MACHINE_KEY, get_admin_owner_machine_id(connection)])
    meta_sheet.append(["last_exported_at", auth_now()])

    try:
        workbook.save(ACCOUNT_SHEET_PATH)
    except PermissionError as exc:
        if not raise_on_error:
            return
        raise HTTPException(
            status_code=409,
            detail=f"Không ghi được file tài khoản. Hãy đóng {ACCOUNT_SHEET_PATH} rồi thử lại.",
        ) from exc


def sync_users_from_sheet_if_needed(connection) -> None:
    """Throttled wrapper — runs at most once every SYNC_THROTTLE_SECONDS per process."""
    global _last_sync_at
    now = time.monotonic()
    with _sync_throttle_lock:
        if now - _last_sync_at < SYNC_THROTTLE_SECONDS:
            return
        _last_sync_at = now
    sync_users_from_sheet(connection)


def sync_users_from_sheet(connection) -> None:
    global _last_sync_at
    _last_sync_at = time.monotonic()
    remote_url = _account_sheet_url()
    if remote_url:
        try:
            rows = _fetch_remote_account_rows(remote_url)
            changed, should_export = _apply_account_rows(
                connection,
                rows,
                allow_admin_roles=False,
            )
            if changed or should_export or not ACCOUNT_SHEET_PATH.exists():
                export_users_to_sheet(connection, raise_on_error=False)
            return
        except (OSError, URLError, TimeoutError, ValueError, csv.Error):
            pass

    sheet_source_path = _account_sheet_source_path()
    if not sheet_source_path.exists():
        if users_exist(connection):
            export_users_to_sheet(connection, raise_on_error=False)
        return

    rows, sheet_meta = _xlsx_account_rows_and_meta(sheet_source_path)
    sheet_owner_machine_id = sheet_meta.get(ADMIN_OWNER_MACHINE_KEY, "")
    if sheet_owner_machine_id:
        set_admin_owner_machine(connection, sheet_owner_machine_id)

    changed, should_export = _apply_account_rows(connection, rows)

    if changed or should_export or sheet_source_path != ACCOUNT_SHEET_PATH:
        export_users_to_sheet(connection, raise_on_error=False)
